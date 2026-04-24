"""
Parser — Module 2 of the ReGenAI pipeline.

Takes a FileEntry from the registry and extracts text content using
the appropriate parser for each file type.  Returns a ParsedFile with
raw text, optional section boundaries, and metadata.

Routing:
    PDF       → unstructured (primary), pdfplumber (fallback)
    DOCX      → unstructured
    DOC       → unstructured
    PPTX      → unstructured (preserves slide boundaries)
    XLSX      → openpyxl (structured sheet extraction)
    Code      → direct read with encoding detection
    Markdown  → direct read
    Config    → direct read
    Other     → attempt direct read, graceful failure
"""

from dataclasses import dataclass, field
from pathlib import Path

import chardet
from rich.console import Console

from regenai.registry import FileEntry

console = Console()


# ---------------------------------------------------------------------------
# Data structures
# ---------------------------------------------------------------------------


@dataclass
class ParsedFile:
    """Result of parsing a single file."""

    entry: FileEntry
    raw_text: str  # full extracted text
    sections: list[str] | None = None  # logical sections (slides, sheets, etc.)
    metadata: dict = field(default_factory=dict)  # parser-extracted metadata
    parse_success: bool = True
    error: str | None = None


# ---------------------------------------------------------------------------
# Encoding detection
# ---------------------------------------------------------------------------


def _read_text_with_encoding(path: Path) -> str:
    """
    Read a text file, trying UTF-8 first, falling back to chardet detection.
    """
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        pass

    # Fallback: detect encoding
    raw_bytes = path.read_bytes()
    detected = chardet.detect(raw_bytes)
    encoding = detected.get("encoding", "utf-8") or "utf-8"

    try:
        return raw_bytes.decode(encoding, errors="replace")
    except (UnicodeDecodeError, LookupError):
        return raw_bytes.decode("utf-8", errors="replace")


# ---------------------------------------------------------------------------
# Format-specific parsers
# ---------------------------------------------------------------------------


def _parse_pdf(path: Path) -> tuple[str, list[str] | None, dict]:
    """
    Parse PDF using unstructured (handles digital text + images via OCR).
    Falls back to pdfplumber if unstructured fails.
    """
    metadata: dict = {}

    # Primary: unstructured
    try:
        from unstructured.partition.pdf import partition_pdf

        elements = partition_pdf(str(path))

        if elements:
            sections: list[str] = []
            current_section: list[str] = []
            current_page: int | None = None

            for el in elements:
                page = getattr(el.metadata, "page_number", None)
                text = str(el).strip()
                if not text:
                    continue

                # Group by page as sections
                if page != current_page and current_section:
                    sections.append("\n".join(current_section))
                    current_section = []
                current_page = page
                current_section.append(text)

            if current_section:
                sections.append("\n".join(current_section))

            raw_text = "\n\n".join(sections)
            metadata["parser"] = "unstructured"
            metadata["page_count"] = len(sections)

            if raw_text.strip():
                return raw_text, sections, metadata

    except Exception as e:
        metadata["unstructured_error"] = str(e)

    # Fallback: pdfplumber
    try:
        import pdfplumber

        sections = []
        with pdfplumber.open(path) as pdf:
            metadata["page_count"] = len(pdf.pages)
            for page in pdf.pages:
                text = page.extract_text()
                if text and text.strip():
                    sections.append(text.strip())

        raw_text = "\n\n".join(sections)
        metadata["parser"] = "pdfplumber"
        return raw_text, sections if len(sections) > 1 else None, metadata

    except Exception as e:
        raise RuntimeError(f"Both PDF parsers failed: {e}")


def _parse_docx(path: Path) -> tuple[str, list[str] | None, dict]:
    """Parse DOCX using unstructured."""
    from unstructured.partition.docx import partition_docx

    elements = partition_docx(str(path))
    metadata: dict = {"parser": "unstructured"}

    sections: list[str] = []
    current_section: list[str] = []

    for el in elements:
        text = str(el).strip()
        if not text:
            continue

        el_type = type(el).__name__
        # Start a new section on Title or Header elements
        if el_type in ("Title", "Header") and current_section:
            sections.append("\n".join(current_section))
            current_section = []

        current_section.append(text)

    if current_section:
        sections.append("\n".join(current_section))

    raw_text = "\n\n".join(sections)
    return raw_text, sections if len(sections) > 1 else None, metadata


def _parse_doc(path: Path) -> tuple[str, list[str] | None, dict]:
    """Parse old-format DOC using unstructured."""
    from unstructured.partition.doc import partition_doc

    elements = partition_doc(str(path))
    metadata: dict = {"parser": "unstructured"}

    texts = [str(el).strip() for el in elements if str(el).strip()]
    raw_text = "\n\n".join(texts)
    return raw_text, None, metadata


def _parse_pptx(path: Path) -> tuple[str, list[str] | None, dict]:
    """
    Parse PPTX using unstructured.
    Each slide becomes a separate section.
    """
    from unstructured.partition.pptx import partition_pptx

    elements = partition_pptx(str(path))
    metadata: dict = {"parser": "unstructured"}

    sections: list[str] = []
    current_slide: list[str] = []
    current_page: int | None = None

    for el in elements:
        page = getattr(el.metadata, "page_number", None)
        text = str(el).strip()
        if not text:
            continue

        if page != current_page and current_slide:
            sections.append("\n".join(current_slide))
            current_slide = []
        current_page = page
        current_slide.append(text)

    if current_slide:
        sections.append("\n".join(current_slide))

    raw_text = "\n\n".join(sections)
    metadata["slide_count"] = len(sections)
    return raw_text, sections if len(sections) > 1 else None, metadata


def _parse_xlsx(path: Path) -> tuple[str, list[str] | None, dict]:
    """
    Parse XLSX using openpyxl.
    Extracts each sheet as a section with row data as readable text.
    """
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    metadata: dict = {
        "parser": "openpyxl",
        "sheet_names": wb.sheetnames,
        "sheet_count": len(wb.sheetnames),
    }

    sections: list[str] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows: list[str] = [f"Sheet: {sheet_name}"]
        header_row: list[str] = []

        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            # Skip completely empty rows
            values = [str(cell) if cell is not None else "" for cell in row]
            if not any(v.strip() for v in values):
                continue

            if row_idx == 0:
                header_row = values
                rows.append(" | ".join(values))
            else:
                # For data rows with headers, format as "header: value" pairs
                if header_row:
                    pairs = []
                    for h, v in zip(header_row, values):
                        if v.strip():
                            pairs.append(f"{h}: {v}" if h.strip() else v)
                    if pairs:
                        rows.append(", ".join(pairs))
                else:
                    rows.append(" | ".join(values))

        if len(rows) > 1:  # more than just the sheet name
            sections.append("\n".join(rows))

    wb.close()

    raw_text = "\n\n".join(sections)
    return raw_text, sections if len(sections) > 1 else None, metadata


def _parse_tex(path: Path) -> tuple[str, list[str] | None, dict]:
    """
    Parse LaTeX files — read directly but extract section structure.
    """
    import re

    text = _read_text_with_encoding(path)
    metadata: dict = {"parser": "direct"}

    # Split on LaTeX section commands for section boundaries
    section_pattern = re.compile(r"\\(?:section|subsection|chapter)\{", re.IGNORECASE)

    parts = section_pattern.split(text)
    if len(parts) > 1:
        return text, [p.strip() for p in parts if p.strip()], metadata

    return text, None, metadata


def _parse_csv_tsv(path: Path, delimiter: str) -> tuple[str, list[str] | None, dict]:
    """Parse CSV/TSV into readable text."""
    import csv

    text = _read_text_with_encoding(path)
    metadata: dict = {"parser": "csv"}

    try:
        reader = csv.reader(text.splitlines(), delimiter=delimiter)
        rows = list(reader)
        if not rows:
            return text, None, metadata

        header = rows[0]
        metadata["columns"] = header
        metadata["row_count"] = len(rows) - 1

        lines: list[str] = [" | ".join(header)]
        for row in rows[1:]:
            pairs = []
            for h, v in zip(header, row):
                if v.strip():
                    pairs.append(f"{h}: {v}")
            if pairs:
                lines.append(", ".join(pairs))

        return "\n".join(lines), None, metadata
    except csv.Error:
        # Fallback to raw text
        return text, None, metadata


# ---------------------------------------------------------------------------
# Main parse dispatcher
# ---------------------------------------------------------------------------


def parse_file(entry: FileEntry) -> ParsedFile:
    """
    Parse a single file and extract its text content.

    Routes to the appropriate parser based on file type and extension.
    Catches all errors per-file so the pipeline continues on failure.
    """
    try:
        ext = entry.extension.lower()

        # PDF
        if ext == ".pdf":
            raw, sections, meta = _parse_pdf(entry.path)

        # Word documents
        elif ext == ".docx":
            raw, sections, meta = _parse_docx(entry.path)
        elif ext == ".doc":
            raw, sections, meta = _parse_doc(entry.path)

        # Presentations
        elif ext in (".pptx", ".ppt"):
            raw, sections, meta = _parse_pptx(entry.path)

        # Spreadsheets
        elif ext in (".xlsx", ".xls", ".ods"):
            raw, sections, meta = _parse_xlsx(entry.path)
        elif ext == ".csv":
            raw, sections, meta = _parse_csv_tsv(entry.path, delimiter=",")
        elif ext == ".tsv":
            raw, sections, meta = _parse_csv_tsv(entry.path, delimiter="\t")

        # LaTeX
        elif ext == ".tex":
            raw, sections, meta = _parse_tex(entry.path)

        # Code, markdown, text, config — direct read
        elif entry.file_type in ("code", "config", "data") or ext in (
            ".md",
            ".txt",
            ".rst",
            ".html",
            ".htm",
            ".json",
            ".yaml",
            ".yml",
            ".toml",
            ".ini",
            ".cls",
            ".bib",
            ".bst",
            ".sty",
        ):
            raw = _read_text_with_encoding(entry.path)
            sections = None
            meta = {"parser": "direct"}

        # Unknown — attempt direct read
        else:
            raw = _read_text_with_encoding(entry.path)
            sections = None
            meta = {"parser": "direct_fallback"}

        # Check if we got usable content
        if not raw or not raw.strip():
            return ParsedFile(
                entry=entry,
                raw_text="",
                sections=None,
                metadata=meta,
                parse_success=False,
                error="No text content extracted",
            )

        return ParsedFile(
            entry=entry,
            raw_text=raw,
            sections=sections,
            metadata=meta,
        )

    except Exception as e:
        console.print(f"  [red]Parse error:[/red] {entry.relative_path}: {e}")
        return ParsedFile(
            entry=entry,
            raw_text="",
            parse_success=False,
            error=str(e),
        )


# ---------------------------------------------------------------------------
# Batch parsing with progress
# ---------------------------------------------------------------------------


def parse_all(entries: list[FileEntry]) -> list[ParsedFile]:
    """
    Parse all files in the registry with a progress indicator.
    Returns list of ParsedFile objects (including failures).
    """
    from rich.progress import Progress, SpinnerColumn, TextColumn, BarColumn

    results: list[ParsedFile] = []
    success = 0
    failed = 0

    with Progress(
        SpinnerColumn(),
        TextColumn("[progress.description]{task.description}"),
        BarColumn(),
        TextColumn("{task.completed}/{task.total}"),
        console=console,
    ) as progress:
        task = progress.add_task("Parsing files", total=len(entries))

        for entry in entries:
            parsed = parse_file(entry)
            results.append(parsed)

            if parsed.parse_success:
                success += 1
            else:
                failed += 1

            progress.advance(task)

    console.print(f"  [green]Parsed successfully:[/green] {success}")
    if failed:
        console.print(f"  [yellow]Failed to parse:[/yellow] {failed}")
    console.print()

    return results
